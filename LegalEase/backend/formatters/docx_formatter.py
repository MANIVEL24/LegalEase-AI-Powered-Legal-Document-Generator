from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.formatters.common import sanitize_text

BASE_DIR = Path(__file__).resolve().parents[2]
LOGO = BASE_DIR / "assets" / "logo.png"

def _set_cell_text(cell, text):
    cell.text = text
    for run in cell.paragraphs[0].runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)

def format_docx(text: str, doc_type: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    if LOGO.exists():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(1.15))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)

    lines = sanitize_text(text).splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            document.add_paragraph("")
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        if stripped.startswith("DRAFT"):
            r = p.add_run(stripped)
            r.bold = True
        elif stripped.isupper() and len(stripped) < 100:
            r = p.add_run(stripped)
            r.bold = True
        else:
            p.add_run(stripped)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-generated draft. Review before use.").font.size = Pt(8)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
